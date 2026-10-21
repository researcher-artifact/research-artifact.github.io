#!/usr/bin/env python3
"""Add the RQ5 classification and the technology readiness assessment to DEF_TA.xlsx.

The RQ5 approach groups are derived from the final codes of the thematic analysis and from
technical terms recorded in the data extraction form, using the rules in
analysis_definitions.json (Table approach_groups of the article). The TRL of each solution is
an input: it is read from the per-solution ratings workbook (sheet "Solutions").

The sheets are written straight into the workbook package, so the existing sheets, charts,
tables and data model are kept as they are. Run it on a workbook that does not have the sheets
yet, or it stops. Existing cells changed: "RQ5. Approaches" (counts, names and references of
the groups) and "ApproachesAnalyses" (groups of each study).

Usage:
    python3 add_analysis_sheets.py --workbook ../data/DEF_TA.xlsx --tra TRA_solutions.xlsx \
        [--previous RQ5_approach_groups.xlsx] --out ../data/DEF_TA.xlsx
"""
import argparse
import json
import os
import re
import shutil
import statistics
import tempfile
import zipfile
from collections import Counter
from xml.sax.saxutils import escape

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
NS_REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
WS_TYPE = NS_REL + '/worksheet'
WS_CT = 'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'
TRLS = [2, 3, 4, 5, 7]  # levels with at least one solution; checked below


# ------------------------------------------------------------------ reading

def text(v):
    return '' if v is None else str(v).strip()


def rows_of(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [text(h) for h in rows[0]], rows[1:]


def col(header, pattern):
    rx = re.compile(pattern, re.I)
    for i, h in enumerate(header):
        if rx.search(h):
            return i
    raise KeyError(pattern)


def vocabulary(wb, sheet):
    out = []
    for r in list(wb[sheet].iter_rows(values_only=True))[1:]:
        if not r[0]:
            break
        try:
            float(r[1])
            out.append(text(r[0]))
        except (TypeError, ValueError):
            pass
    return out


def split(value, vocab):
    """Splits a coded cell into category names, known names first (some contain commas)."""
    rest, found = text(value), []
    if rest in ('', '0'):
        return found
    for name in sorted(vocab, key=len, reverse=True):
        rx = re.compile(re.escape(name), re.I)
        if rx.search(rest):
            found.append(name)
            rest = rx.sub(' ', rest)
    for s in re.split(r'[,;\n]+', rest):
        if s.strip() and s.strip() != '0':
            raise ValueError(f'Category not in the summary sheet: {s.strip()!r}')
    return found


def paper_key(pid):
    m = re.search(r'\d+', pid or '')
    return int(m.group(0)) if m else 10 ** 6


# ------------------------------------------------------------- RQ5 rules

def find_terms(de_row, de_header, defs):
    """Technical terms found in the DEF fields listed in the definitions, with the fields."""
    found = []
    for t in defs['terms']:
        fields = []
        for name in defs['term_fields']:
            value = text(de_row[de_header.index(name)]).lower()
            if re.search(t['pattern'], value) and (not t.get('also_requires') or re.search(t['also_requires'], value)):
                fields.append(name)
        # "decentralised identifier + identity": both may sit in different fields of the same study
        if not fields and t.get('also_requires'):
            joined = ' '.join(text(de_row[de_header.index(n)]).lower() for n in defs['term_fields'])
            if re.search(t['pattern'], joined) and re.search(t['also_requires'], joined):
                fields = [n for n in defs['term_fields'] if re.search(t['pattern'], text(de_row[de_header.index(n)]).lower())]
        if fields:
            found.append((t['term'], fields))
    return found


def clause_evidence(clause, T, P, O, terms):
    """Returns the codes that satisfy one clause, or None when it is not satisfied."""
    ev = []
    for key, have in (('technologies', T), ('purposes', P), ('objects', O), ('terms', terms)):
        need_all = clause.get(f'{key}_all')
        need_any = clause.get(f'{key}_any')
        if need_all:
            if not all(x in have for x in need_all):
                return None
            ev += need_all
        if need_any:
            hit = [x for x in need_any if x in have]
            if not hit:
                return None
            ev += hit
    return ev


def classify(T, P, O, terms, defs):
    out = {}
    for g in defs['groups']:
        for clause in g['clauses']:
            ev = clause_evidence(clause, T, P, O, terms)
            if ev is not None:
                out[g['code']] = ev
                break
    return out


def short_field(name):
    m = re.match(r'(.*?)\s*\(AI Tool - (?:OpenAI |Anthropic )?(\w+)\)', name)
    base = 'Technology Stack' if name.startswith('Technology') else 'Proposed Solution'
    return f'{base} ({m.group(2)})' if m else name


# ------------------------------------------------------------- sheet XML

def cell_ref(c, r):
    s = ''
    c += 1
    while c:
        c, rem = divmod(c - 1, 26)
        s = chr(65 + rem) + s
    return f'{s}{r}'


ILLEGAL = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')


def sheet_xml(rows, widths, styles, header_rows=(0,), percent_cols=()):
    """rows: list of lists (str, int, float or None). Header rows use the header style."""
    out = []
    ncols = max(len(r) for r in rows)
    for ri, row in enumerate(rows, start=1):
        cells = []
        for ci, v in enumerate(row):
            if v is None or v == '':
                continue
            ref = cell_ref(ci, ri)
            if ri - 1 in header_rows:
                s = styles['header']
            elif ci in percent_cols and isinstance(v, float):
                s = styles['percent']
            else:
                s = styles['body']
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                cells.append(f'<c r="{ref}" s="{s}"><v>{v}</v></c>')
            else:
                t = escape(ILLEGAL.sub('', str(v)))
                cells.append(f'<c r="{ref}" s="{s}" t="inlineStr"><is><t xml:space="preserve">{t}</t></is></c>')
        out.append(f'<row r="{ri}">{"".join(cells)}</row>')
    cols = ''.join(f'<col min="{i + 1}" max="{i + 1}" width="{w}" customWidth="1"/>' for i, w in enumerate(widths))
    dim = f'A1:{cell_ref(ncols - 1, len(rows))}'
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            f'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="{NS_REL}">'
            f'<dimension ref="{dim}"/><sheetViews><sheetView workbookViewId="0">'
            '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/></sheetView></sheetViews>'
            f'<sheetFormatPr defaultRowHeight="12.5"/><cols>{cols}</cols><sheetData>{"".join(out)}</sheetData>'
            '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>')


def add_styles(styles_xml):
    """Appends a header, a body and a percentage cell style; returns the new XML and their indices."""
    m = re.search(r'<cellXfs count="(\d+)">', styles_xml)
    n = int(m.group(1))
    bold = re.findall(r'<font>.*?</font>|<font/>', re.search(r'<fonts[^>]*>(.*?)</fonts>', styles_xml, re.S).group(1), re.S)
    header_font = next(i for i, f in enumerate(bold) if '<b/>' in f and 'FFFFFFFF' in f)
    new = (f'<xf numFmtId="0" fontId="{header_font}" fillId="4" borderId="0" xfId="0" applyFont="1" applyFill="1" applyAlignment="1">'
           '<alignment vertical="center" wrapText="1"/></xf>'
           '<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0" applyAlignment="1"><alignment vertical="top" wrapText="1"/></xf>'
           '<xf numFmtId="10" fontId="0" fillId="0" borderId="0" xfId="0" applyNumberFormat="1" applyAlignment="1"><alignment vertical="top"/></xf>')
    styles_xml = styles_xml.replace(m.group(0), f'<cellXfs count="{n + 3}">', 1)
    end = styles_xml.index('</cellXfs>')
    styles_xml = styles_xml[:end] + new + styles_xml[end:]
    return styles_xml, {'header': n, 'body': n + 1, 'percent': n + 2}


def set_cell(sheet_xml_text, ref, value, keep_formula=False):
    """Replaces the value of one existing cell, keeping its style (and formula when asked)."""
    m = re.search(rf'<c r="{ref}"(?: [^>]*)?(?:/>|>.*?</c>)', sheet_xml_text, re.S)
    if not m:
        raise KeyError(ref)
    old = m.group(0)
    style = re.search(r' s="(\d+)"', old)
    s = f' s="{style.group(1)}"' if style else ''
    if keep_formula:
        f = re.search(r'<f>.*?</f>|<f [^>]*/>|<f [^>]*>.*?</f>', old, re.S).group(0)
        if isinstance(value, (int, float)):
            new = f'<c r="{ref}"{s}>{f}<v>{value}</v></c>'
        else:
            new = f'<c r="{ref}"{s} t="str">{f}<v>{escape(str(value))}</v></c>'
    elif value is None or value == '':
        new = f'<c r="{ref}"{s}/>'
    elif isinstance(value, (int, float)):
        new = f'<c r="{ref}"{s}><v>{value}</v></c>'
    else:
        new = f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{escape(str(value))}</t></is></c>'
    return sheet_xml_text.replace(old, new, 1)


# ------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--workbook', required=True)
    ap.add_argument('--tra', required=True, help='workbook with the TRL of each solution (sheet "Solutions")')
    ap.add_argument('--previous', help='workbook with the previous RQ5 assignment (sheet "Assignments"), for the "Changed" column')
    ap.add_argument('--definitions', default=os.path.join(HERE, 'analysis_definitions.json'))
    ap.add_argument('--out', required=True)
    args = ap.parse_args()
    defs = json.load(open(args.definitions, encoding='utf-8'))
    groups = defs['groups']
    gname = {g['code']: g['name'] for g in groups}
    codes = [g['code'] for g in groups]

    wb = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    if any(n.startswith(('RQ5 Rules', 'TRA Solutions')) for n in wb.sheetnames):
        raise SystemExit('The workbook already has the analysis sheets.')
    V = {k: vocabulary(wb, s) for k, s in (('T', 'RQ2. Technologies Summary'), ('P', 'RQ3. Purposes Summary'), ('O', 'RQ1. Digital Objects Summary'))}
    ta_h, ta = rows_of(wb['Thematic Analysis'])
    de_h, de = rows_of(wb['Data Extraction'])
    de_by = {text(r[col(de_h, r'^r[a]?y+an id$')]): r for r in de}
    iY = col(de_h, r'^year$')

    previous = {}
    if args.previous:
        pw = openpyxl.load_workbook(args.previous, read_only=True)
        ph, pr = rows_of(pw['Assignments'])
        for r in pr:
            previous[text(r[0])] = sorted(set(re.findall(r'AC-\d', text(r[col(ph, r'^previous')]))))

    studies = []
    for r in ta:
        rid = text(r[col(ta_h, r'^r[a]?y+an id$')])
        if not rid:
            continue
        d = de_by[rid]
        T = split(r[col(ta_h, r'^technolog')], V['T'])
        P = split(r[col(ta_h, r'^purposes?$')], V['P'])
        O = split(r[col(ta_h, r'^digital objects?$')], V['O'])
        terms = find_terms(d, de_h, defs)
        ev = classify(T, P, O, [t for t, _ in terms], defs)
        studies.append({
            'article': text(r[col(ta_h, r'^article id$')]), 'paper': text(d[col(de_h, r'^pap+er$')]), 'id': rid,
            'year': int(text(d[iY])), 'T': T, 'P': P, 'O': O, 'terms': terms, 'ev': ev,
            'groups': [c for c in codes if c in ev],
        })
    studies.sort(key=lambda s: paper_key(s['paper']))
    n = len(studies)
    by_id = {s['id']: s for s in studies}

    # TRA ratings
    tw = openpyxl.load_workbook(args.tra, read_only=True)
    th, tr = rows_of(tw['Solutions'])
    tra = {}
    for r in tr:
        rid = text(r[col(th, r'^citation key|rayyan')])
        if not rid:
            continue
        tra[rid] = {
            'trl': int(r[col(th, r'^solution trl')]), 'env': text(r[col(th, r'environment')]),
            'criterion': text(r[col(th, r'^criterion')]), 'source': text(r[col(th, r'^evidence source')]),
            'evidence': text(r[col(th, r'^evidence$')]),
        }
    missing = set(by_id) ^ set(tra)
    if missing:
        raise SystemExit(f'Studies and TRA ratings differ: {sorted(missing)}')
    levels = sorted({t['trl'] for t in tra.values()})
    if levels != TRLS:
        raise SystemExit(f'TRL levels {levels} differ from {TRLS}; update TRLS.')
    for s in studies:
        s['tra'] = tra[s['id']]

    def period(y):
        for a, b in defs['periods']:
            if a <= y <= b:
                return f'{a}–{b}'
        raise ValueError(y)

    refs = lambda lst: ', '.join(sorted(s['id'] for s in lst))
    pct = lambda k, total: round(k / total, 6) if total else 0.0
    sheets = []

    # RQ5 Rules: one row per clause; clauses of a group are alternatives (OR), conditions in a row all hold (AND).
    rows = [['Code', 'Approach group', 'Definition', 'Attribution rule', 'Clause', 'Technologies (all of)', 'Technologies (any of)',
             'Purposes (any of)', 'Digital objects (any of)', 'Technical terms (any of)']]
    for g in groups:
        for i, c in enumerate(g['clauses'], start=1):
            rows.append([g['code'], g['name'], g['definition'] if i == 1 else '', g['rule'] if i == 1 else '', i,
                         '; '.join(c.get('technologies_all', [])), '; '.join(c.get('technologies_any', [])),
                         '; '.join(c.get('purposes_any', [])), '; '.join(c.get('objects_any', [])), '; '.join(c.get('terms_any', []))])
    sheets.append(('RQ5 Rules', rows, [8, 30, 70, 40, 7, 26, 26, 34, 18, 34]))

    rows = [['Technical term', 'Pattern (regular expression, case-insensitive)', 'Also requires', 'Fields searched']]
    for t in defs['terms']:
        rows.append([t['term'], t['pattern'], t.get('also_requires', ''), '; '.join(defs['term_fields'])])
    sheets.append(('RQ5 Terms', rows, [32, 34, 14, 90]))

    # RQ5 Classification: one row per study.
    rows = [['Article ID', 'Paper ID', 'Rayyan ID', 'Year', 'Technologies (RQ2)', 'Purposes (RQ3)', 'Digital objects (RQ1)',
             'Technical terms found', 'Fields with the terms'] + [f'{c} {gname[c]}' for c in codes]
            + ['Approach groups', 'No. of groups', 'Previous approach groups', 'Changed']]
    for s in studies:
        prev = previous.get(s['id'])
        rows.append([s['article'], s['paper'], s['id'], s['year'], '; '.join(s['T']), '; '.join(s['P']), '; '.join(s['O']),
                     '; '.join(t for t, _ in s['terms']),
                     '; '.join(f"{t}: {', '.join(short_field(f) for f in fs)}" for t, fs in s['terms'])]
                    + ['; '.join(s['ev'][c]) if c in s['ev'] else '' for c in codes]
                    + [', '.join(s['groups']), len(s['groups']),
                       ', '.join(c for c in codes if c in (prev or [])) if prev is not None else '',
                       ('Yes' if sorted(s['groups']) != prev else 'No') if prev is not None else ''])
    sheets.append(('RQ5 Classification', rows, [20, 9, 17, 6, 34, 34, 28, 22, 30] + [20] * len(codes) + [22, 8, 22, 8]))

    # RQ5 Distribution
    count = {c: [s for s in studies if c in s['groups']] for c in codes}
    prev_count = Counter(c for p in previous.values() for c in p)
    rows = [['Code', 'Approach group', 'No. of studies', '% of studies', 'Previous no.', 'References']]
    for c in sorted(codes, key=lambda c: -len(count[c])):
        rows.append([c, gname[c], len(count[c]), pct(len(count[c]), n), prev_count.get(c, '') if previous else '', refs(count[c])])
    none = [s for s in studies if not s['groups']]
    rows.append(['–', 'No group', len(none), pct(len(none), n), '', refs(none)])
    sheets.append(('RQ5 Distribution', rows, [8, 50, 10, 10, 10, 120]))

    # RQ5 Combinations
    combos = Counter(', '.join(s['groups']) for s in studies)
    rows = [['Combination of groups', 'No. of groups', 'No. of studies', '% of studies', 'References']]
    for combo, k in sorted(combos.items(), key=lambda kv: (-kv[1], kv[0])):
        lst = [s for s in studies if ', '.join(s['groups']) == combo]
        rows.append([combo or 'No group', len(lst[0]['groups']), k, pct(k, n), refs(lst)])
    sheets.append(('RQ5 Combinations', rows, [34, 10, 10, 10, 120]))

    # RQ5 Statistics
    dom = defs['dominant']
    multi = [s for s in studies if len(s['groups']) >= 2]
    single = [s for s in studies if len(s['groups']) == 1]
    any_dom = [s for s in studies if set(s['groups']) & set(dom)]
    only_dom = [s for s in any_dom if set(s['groups']) <= set(dom)]
    exact = [s for s in studies if set(s['groups']) == set(dom)]
    all_dom = [s for s in studies if set(dom) <= set(s['groups'])]
    fa = [s for s in studies if 'Forensic Attribution & Post-Incident Investigation' in s['P']]
    fa3 = [s for s in fa if 'AC-3' in s['groups']]
    assignments = sum(len(s['groups']) for s in studies)
    rows = [['Metric', 'Value', '% of studies', 'References'],
            ['Studies', n, '', ''],
            ['Group assignments', assignments, '', ''],
            ['Mean groups per study', round(assignments / n, 2), '', ''],
            ['Studies in two or more groups', len(multi), pct(len(multi), n), refs(multi)],
            ['Studies in exactly one group', len(single), pct(len(single), n), refs(single)],
            ['Studies in no group', len(none), pct(len(none), n), refs(none)],
            [f'Studies in at least one dominant group ({", ".join(dom)})', len(any_dom), pct(len(any_dom), n), refs(any_dom)],
            ['Studies only in dominant groups', len(only_dom), pct(len(only_dom), n), refs(only_dom)],
            ['Studies in exactly the three dominant groups', len(exact), pct(len(exact), n), refs(exact)],
            ['Studies in all three dominant groups (possibly with others)', len(all_dom), pct(len(all_dom), n), refs(all_dom)],
            ['Studies with purpose Forensic Attribution & Post-Incident Investigation', len(fa), pct(len(fa), n), refs(fa)],
            ['… of which in AC-3 Embedded Forensic Attribution', len(fa3), pct(len(fa3), n), refs(fa3)]]
    sheets.append(('RQ5 Statistics', rows, [60, 10, 10, 120]))

    # TRL Rubric
    rows = [['TRL', 'NASA definition (NASA/SP-20205003605)', 'Criterion applied to the mapped studies', 'Environment']]
    rows += [[r['trl'], r['definition'], r['criterion'], r['environment']] for r in defs['trl_rubric']]
    rows += [[], ['Rule', defs['trl_rule']]]
    sheets.append(('TRL Rubric', rows, [6, 50, 80, 22]))

    # TRA Solutions: one row per solution.
    rows = [['Article ID', 'Paper ID', 'Rayyan ID', 'Year', 'Period', 'Technologies (RQ2)', 'Approach groups (RQ5)', 'TRL',
             'Validation environment', 'Criterion met', 'Evidence source', 'Evidence']]
    for s in studies:
        t = s['tra']
        rows.append([s['article'], s['paper'], s['id'], s['year'], period(s['year']), '; '.join(s['T']), ', '.join(s['groups']),
                     t['trl'], t['env'], t['criterion'], t['source'], t['evidence']])
    sheets.append(('TRA Solutions', rows, [20, 9, 17, 6, 10, 34, 20, 5, 18, 50, 14, 80]))

    def trl_row(label, lst):
        trl = [s['tra']['trl'] for s in lst]
        return [label, len(lst)] + [trl.count(k) for k in TRLS] + [statistics.median_low(trl) if trl else '', max(trl) if trl else '']

    # TRL Distribution
    rows = [['TRL', 'NASA definition', 'No. of solutions', '% of solutions', 'References']]
    for r in defs['trl_rubric']:
        lst = [s for s in studies if s['tra']['trl'] == r['trl']]
        rows.append([r['trl'], r['definition'], len(lst), pct(len(lst), n), refs(lst)])
    sheets.append(('TRL Distribution', rows, [6, 50, 10, 10, 120]))

    head = ['No. of solutions'] + [f'TRL {k}' for k in TRLS] + ['Median TRL', 'Maximum TRL']
    tech_count = Counter(t for s in studies for t in s['T'])
    rows = [['Technology'] + head]
    for t in sorted(tech_count, key=lambda t: (-tech_count[t], t)):
        rows.append(trl_row(t, [s for s in studies if t in s['T']]))
    rows.append(trl_row('All solutions', studies))
    sheets.append(('TRL Technologies', rows, [40, 10] + [7] * len(TRLS) + [10, 10]))

    rows = [['Approach group'] + head]
    for c in sorted(codes, key=lambda c: -len(count[c])):
        rows.append(trl_row(f'{c} {gname[c]}', count[c]))
    rows.append(trl_row('No group', none))
    rows.append(trl_row('All solutions', studies))
    sheets.append(('TRL Approaches', rows, [60, 10] + [7] * len(TRLS) + [10, 10]))

    rows = [['Period'] + head + ['Solutions at TRL 4 or above', '% at TRL 4 or above']]
    for a, b in defs['periods']:
        lst = [s for s in studies if a <= s['year'] <= b]
        up = [s for s in lst if s['tra']['trl'] >= 4]
        rows.append(trl_row(f'{a}–{b}', lst) + [len(up), pct(len(up), len(lst))])
    up = [s for s in studies if s['tra']['trl'] >= 4]
    rows.append(trl_row('All solutions', studies) + [len(up), pct(len(up), n)])
    sheets.append(('TRL Periods', rows, [14, 10] + [7] * len(TRLS) + [10, 10, 12, 12]))

    percent_cols = {'RQ5 Distribution': (3,), 'RQ5 Combinations': (3,), 'RQ5 Statistics': (2,), 'TRL Distribution': (3,),
                    'TRL Periods': (len(head) + 2,)}

    # --------------------------------------------------------- write package
    tmp = tempfile.mkdtemp()
    try:
        with zipfile.ZipFile(args.workbook) as z:
            z.extractall(tmp)
            order = z.namelist()
            # Keep each part's compression; the data model (xl/model/item.data) is stored uncompressed.
            kind = {i.filename: i.compress_type for i in z.infolist()}
        p = lambda *a: os.path.join(tmp, *a)
        read = lambda f: open(p(f), encoding='utf-8').read()
        write = lambda f, s: open(p(f), 'w', encoding='utf-8').write(s)

        styles, st = add_styles(read('xl/styles.xml'))
        write('xl/styles.xml', styles)
        wbx = read('xl/workbook.xml')
        rels = read('xl/_rels/workbook.xml.rels')
        ct = read('[Content_Types].xml')
        next_rid = max(int(x) for x in re.findall(r'Id="rId(\d+)"', rels)) + 1
        next_sid = max(int(x) for x in re.findall(r'sheetId="(\d+)"', wbx)) + 1
        next_file = max(int(x) for x in re.findall(r'worksheets/sheet(\d+)\.xml', rels)) + 1
        new_sheets, new_files = '', []
        for name, rows, widths in sheets:
            f = f'xl/worksheets/sheet{next_file}.xml'
            write(f, sheet_xml(rows, widths, st, percent_cols=percent_cols.get(name, ())))
            new_files.append(f)
            new_sheets += f'<sheet name="{escape(name)}" sheetId="{next_sid}" r:id="rId{next_rid}"/>'
            rels = rels.replace('</Relationships>', f'<Relationship Id="rId{next_rid}" Type="{WS_TYPE}" Target="worksheets/sheet{next_file}.xml"/></Relationships>')
            ct = ct.replace('</Types>', f'<Override PartName="/{f}" ContentType="{WS_CT}"/></Types>')
            next_rid, next_sid, next_file = next_rid + 1, next_sid + 1, next_file + 1
        wbx = wbx.replace('</sheets>', new_sheets + '</sheets>')
        # Recalculate formulas when the workbook is opened, since the cached values of edited cells were rewritten by hand.
        if '<calcPr' in wbx:
            wbx = re.sub(r'<calcPr([^>]*?)(/?)>', lambda m: '<calcPr' + re.sub(r'\s*fullCalcOnLoad="[^"]*"', '', m.group(1)) + ' fullCalcOnLoad="1"' + m.group(2) + '>', wbx, count=1)
        else:
            wbx = wbx.replace('</workbook>', '<calcPr fullCalcOnLoad="1"/></workbook>')
        write('xl/workbook.xml', wbx)
        write('xl/_rels/workbook.xml.rels', rels)
        write('[Content_Types].xml', ct)

        app = read('docProps/app.xml')
        m = re.search(r'<vt:i4>(\d+)</vt:i4>', app)
        total = int(m.group(1)) + len(sheets)
        app = app.replace(m.group(0), f'<vt:i4>{total}</vt:i4>', 1)
        app = re.sub(r'<TitlesOfParts><vt:vector size="\d+"', f'<TitlesOfParts><vt:vector size="{total}"', app)
        app = app.replace('</vt:vector></TitlesOfParts>', ''.join(f'<vt:lpstr>{escape(s[0])}</vt:lpstr>' for s in sheets) + '</vt:vector></TitlesOfParts>')
        write('docProps/app.xml', app)

        # Existing sheets: "RQ5. Approaches" and "ApproachesAnalyses".
        target = lambda name: 'xl/' + re.search(
            r'Id="%s"[^>]*Target="([^"]+)"' % re.search(r'<sheet name="%s" sheetId="\d+" r:id="(rId\d+)"' % re.escape(escape(name)), wbx).group(1), rels).group(1)
        f7 = target('RQ5. Approaches')
        x = read(f7)
        ordered = sorted(codes, key=lambda c: -len(count[c]))
        for i, c in enumerate(ordered, start=2):
            k = len(count[c])
            share = f'{100 * k / n:.2f}'
            ref = refs(count[c])
            x = set_cell(x, f'A{i}', c)
            x = set_cell(x, f'B{i}', gname[c])
            x = set_cell(x, f'C{i}', k)
            x = set_cell(x, f'D{i}', share.replace('.', ',') + '%')
            x = set_cell(x, f'E{i}', ref)
            short = f'{c} & ${k}$ & ${share.replace(".", "{.}")}\\%$'
            x = set_cell(x, f'F{i}', short + ' \\\\', keep_formula=True)
            x = set_cell(x, f'G{i}', short + f' \\cite{{{ref}}} \\\\', keep_formula=True)
        x = set_cell(x, 'B13', n)
        x = set_cell(x, 'B14', assignments, keep_formula=True)
        x = set_cell(x, 'B15', round(assignments / n, 2), keep_formula=True)
        x = set_cell(x, 'B16', len(multi))
        x = set_cell(x, 'B17', len(single))
        x = set_cell(x, 'B18', ordered[0], keep_formula=True)
        x = set_cell(x, 'B19', ordered[-1], keep_formula=True)
        x = set_cell(x, 'B20', len(none))
        write(f7, x)

        f16 = target('ApproachesAnalyses')
        x = read(f16)
        for r in re.findall(r'<row r="(\d+)"', x):
            if r == '1':
                continue
            a = re.search(rf'<c r="A{r}"[^>]*?(?:/>|>(.*?)</c>)', x, re.S)
            art = re.sub('<[^>]+>', '', a.group(1) or '')
            if a and 't="s"' in a.group(0):
                art = shared_string(tmp, int(art))
            s = next((s for s in studies if s['article'] == art), None)
            if s is None:
                raise SystemExit(f'ApproachesAnalyses row {r}: {art!r} is not a study')
            x = set_cell(x, f'D{r}', ', '.join(s['groups']))
        write(f16, x)

        # xl/calcChain.xml is left as it is: formula cells keep their formulas, so the calculation chain stays valid.
        out_tmp = args.out + '.tmp'
        with zipfile.ZipFile(out_tmp, 'w', zipfile.ZIP_DEFLATED) as z:
            for name in order + new_files:
                z.write(p(name), name, kind.get(name, zipfile.ZIP_DEFLATED))
        os.replace(out_tmp, args.out)
    finally:
        shutil.rmtree(tmp)

    print(f'{len(sheets)} sheets added to {args.out}')
    print('RQ5:', ', '.join(f'{c} {len(count[c])}' for c in ordered), f'· no group {len(none)} · assignments {assignments}')
    print('TRL:', ', '.join(f'{k}: {sum(1 for s in studies if s["tra"]["trl"] == k)}' for k in TRLS))
    if previous:
        print('Changed:', sum(1 for s in studies if sorted(s['groups']) != previous.get(s['id'])))


_SHARED = {}


def shared_string(root, i):
    if root not in _SHARED:
        xml = open(os.path.join(root, 'xl/sharedStrings.xml'), encoding='utf-8').read()
        _SHARED[root] = [re.sub('<[^>]+>', '', s) for s in re.findall(r'<si>(.*?)</si>', xml, re.S)]
    from xml.sax.saxutils import unescape
    return unescape(_SHARED[root][i])


if __name__ == '__main__':
    main()
